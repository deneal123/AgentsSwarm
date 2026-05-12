"""PPTX generation tool: creates PowerPoint presentations from LLM-structured content."""

import io
import json
import logging

from service.services.agents.domain.client import create_chat_completion

logger = logging.getLogger(__name__)

_PLAN_PROMPT = """Ты — ведущий консультант по бизнес-презентациям.
Нужно подготовить структуру PPTX по теме: {topic}

Цель: сделать убедительную и логичную презентацию, где каждый слайд несёт ценность.

Требования к структуре:
- 6-12 слайдов,
- первый слайд: title + subtitle,
- далее: content-слайды с ясным заголовком и 3-6 сильными bullets,
- финальный слайд: выводы / следующие шаги / призыв к действию.

Требования к bullets:
- короткие, конкретные, без общих фраз,
- ориентированы на решения, метрики, риски, действия,
- без дублирования между слайдами.

Верни ТОЛЬКО JSON, без markdown и комментариев:
{{
    "title": "Название презентации",
    "slides": [
        {{
            "type": "title",
            "title": "...",
            "subtitle": "..."
        }},
        {{
            "type": "content",
            "title": "Заголовок слайда",
            "bullets": ["тезис 1", "тезис 2", "тезис 3"]
        }}
    ]
}}"""


def _build_pptx(structure: dict) -> bytes:
    """Build a PPTX file from structured content dict. Returns bytes."""
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    DARK_BG = RGBColor(0x1A, 0x1A, 0x2E)
    ACCENT = RGBColor(0x16, 0x21, 0x3E)
    WHITE = RGBColor(0xFF, 0xFF, 0xFF)
    LIGHT_BLUE = RGBColor(0x4A, 0x9E, 0xD6)

    def _set_bg(slide, color: RGBColor):
        fill = slide.background.fill
        fill.solid()
        fill.fore_color.rgb = color

    def _add_textbox(
        slide,
        text,
        left,
        top,
        width,
        height,
        size=18,
        bold=False,
        color=WHITE,
        align=PP_ALIGN.LEFT,
        wrap=True,
    ):
        tx_box = slide.shapes.add_textbox(left, top, width, height)
        tf = tx_box.text_frame
        tf.word_wrap = wrap
        p = tf.paragraphs[0]
        p.alignment = align
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color
        return tx_box

    slides_data = structure.get("slides", [])

    for i, slide_data in enumerate(slides_data):
        slide_type = slide_data.get("type", "content")

        layout = prs.slide_layouts[6]
        slide = prs.slides.add_slide(layout)
        _set_bg(slide, DARK_BG)

        if slide_type == "title":
            shape = slide.shapes.add_shape(1, Inches(0), Inches(3.2), Inches(13.33), Inches(0.08))
            shape.fill.solid()
            shape.fill.fore_color.rgb = LIGHT_BLUE
            shape.line.fill.background()

            _add_textbox(
                slide,
                slide_data.get("title", "Презентация"),
                Inches(1),
                Inches(1.5),
                Inches(11.33),
                Inches(1.8),
                size=44,
                bold=True,
                color=WHITE,
                align=PP_ALIGN.CENTER,
            )
            _add_textbox(
                slide,
                slide_data.get("subtitle", ""),
                Inches(1),
                Inches(3.6),
                Inches(11.33),
                Inches(1.2),
                size=24,
                color=LIGHT_BLUE,
                align=PP_ALIGN.CENTER,
            )
        else:
            shape = slide.shapes.add_shape(1, Inches(0), Inches(0), Inches(13.33), Inches(0.08))
            shape.fill.solid()
            shape.fill.fore_color.rgb = LIGHT_BLUE
            shape.line.fill.background()

            _add_textbox(
                slide,
                slide_data.get("title", ""),
                Inches(0.5),
                Inches(0.2),
                Inches(11.5),
                Inches(0.9),
                size=28,
                bold=True,
                color=WHITE,
            )

            sep = slide.shapes.add_shape(1, Inches(0.5), Inches(1.15), Inches(11.5), Inches(0.04))
            sep.fill.solid()
            sep.fill.fore_color.rgb = ACCENT
            sep.line.fill.background()

            bullets = slide_data.get("bullets", [])
            bullet_top = Inches(1.35)
            bullet_height = Inches(0.6)
            for j, bullet in enumerate(bullets[:6]):
                _add_textbox(
                    slide,
                    f"▸  {bullet}",
                    Inches(0.7),
                    bullet_top + j * bullet_height,
                    Inches(11.3),
                    bullet_height,
                    size=18,
                )

        _add_textbox(
            slide,
            f"{i + 1}",
            Inches(12.5),
            Inches(6.8),
            Inches(0.5),
            Inches(0.4),
            size=11,
            color=RGBColor(0x88, 0x88, 0xAA),
            align=PP_ALIGN.RIGHT,
        )

    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


async def generate_pptx(topic: str, model: str) -> tuple[bytes, dict]:
    """Generate a PPTX presentation on the given topic.

    Returns (pptx_bytes, structure_dict).
    """
    try:
        resp = await create_chat_completion(
            messages=[
                {"role": "system", "content": "Ты помощник. Отвечай ТОЛЬКО JSON."},
                {"role": "user", "content": _PLAN_PROMPT.format(topic=topic)},
            ],
            model=model,
            temperature=0.4,
            max_tokens=2000,
        )
        raw = getattr(resp.choices[0].message, "content", "") or "{}"
        raw = raw.strip()
        if raw.startswith("```"):
            raw = raw.split("\n", 1)[-1]
        if raw.endswith("```"):
            raw = raw.rsplit("```", 1)[0]
        structure = json.loads(raw.strip())
    except Exception as exc:
        err_name = exc.__class__.__name__
        if "Timeout" in err_name:
            logger.warning("LLM timeout while planning PPTX structure, using fallback template")
        else:
            logger.warning("Failed to get PPTX structure from LLM (%s), using fallback template", err_name)
        structure = {
            "title": topic,
            "slides": [
                {"type": "title", "title": topic, "subtitle": "Сгенерировано GPTHub AI"},
                {"type": "content", "title": "Введение", "bullets": [f"Тема: {topic}"]},
                {"type": "content", "title": "Выводы", "bullets": ["Спасибо за внимание"]},
            ],
        }

    pptx_bytes = _build_pptx(structure)
    return pptx_bytes, structure
