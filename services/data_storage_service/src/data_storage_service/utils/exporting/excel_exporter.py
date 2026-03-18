"""Create Excel workbooks for exported orders."""

from __future__ import annotations

from io import BytesIO
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from service.utils.exporting.models import OrderExportRow

HEADERS = [
    "Order #",
    "Status",
    "Payment status",
    "Payment method",
    "Delivery type",
    "Subtotal",
    "Promo discount",
    "Bonus discount",
    "Delivery cost",
    "Total",
    "Recipient",
    "Email",
    "Phone",
    "City",
    "Address",
    "Created",
]

CURRENCY_COLUMNS = {6, 7, 8, 9, 10}


def render_orders_excel(orders: Iterable[OrderExportRow]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Orders"
    sheet.append(HEADERS)

    for order in orders:
        sheet.append(
            [
                order.order_number,
                order.status,
                order.payment_status,
                order.payment_method,
                order.delivery_type,
                order.subtotal,
                order.promo_discount,
                order.bonus_discount,
                order.delivery_cost,
                order.total,
                order.recipient_name,
                order.recipient_email or "",
                order.recipient_phone or "",
                order.delivery_city or "",
                order.delivery_address or "",
                order.created_at,
            ]
        )

    for index, width in enumerate(
        [18, 12, 15, 15, 12, 13, 14, 14, 14, 13, 20, 28, 16, 16, 40, 20],
        start=1,
    ):
        sheet.column_dimensions[get_column_letter(index)].width = width

    for cell in sheet[1]:
        cell.alignment = Alignment(horizontal="center")

    for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row):
        for idx in CURRENCY_COLUMNS:
            cell = row[idx - 1]
            cell.number_format = "#,##0.00"
        created_cell = row[15]
        created_cell.number_format = "YYYY-MM-DD HH:MM:SS"

    output = BytesIO()
    workbook.save(output)
    output.seek(0)
    return output.getvalue()
